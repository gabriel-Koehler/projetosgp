"""Leitura e validação de planilhas Excel (.xlsx) e CSV para importação (RF05, RF10 a RF13).

Fluxo: o arquivo é lido, cada linha é validada e o resultado vira uma prévia
(linhas válidas + erros por linha). Nada é gravado aqui: quem grava é o service,
depois que o professor confirma.
"""

import csv
import io
import re
import unicodedata
from dataclasses import dataclass, field

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill

from app.services.errors import ErroDeNegocio

TAMANHO_MAXIMO = 5 * 1024 * 1024  # 5 MB
LINHAS_MAXIMAS = 5000
DIFICULDADES = {"facil": "facil", "media": "media", "dificil": "dificil"}
EMAIL = re.compile(r"^[^@\s]+@[^@\s]+\.[^@\s]+$")


@dataclass
class LinhaImportada:
    linha: int  # número da linha na planilha (cabeçalho = 1)
    dados: dict


@dataclass
class ErroLinha:
    linha: int
    mensagens: list[str]


@dataclass
class Previa:
    validas: list[LinhaImportada] = field(default_factory=list)
    erros: list[ErroLinha] = field(default_factory=list)

    @property
    def total(self) -> int:
        return len(self.validas) + len(self.erros)


def normalizar_texto(valor: str) -> str:
    """Para comparar textos ignorando acentos, maiúsculas e espaços extras."""
    sem_acento = unicodedata.normalize("NFKD", valor).encode("ascii", "ignore").decode()
    return " ".join(sem_acento.casefold().split())


def _chave_coluna(nome: str) -> str:
    return normalizar_texto(str(nome or "")).replace(" ", "_")


def ler_planilha(conteudo: bytes, nome_arquivo: str) -> list[LinhaImportada]:
    if len(conteudo) > TAMANHO_MAXIMO:
        raise ErroDeNegocio("Arquivo muito grande (máximo 5 MB).")
    if not conteudo:
        raise ErroDeNegocio("O arquivo está vazio.")

    extensao = nome_arquivo.lower().rsplit(".", 1)[-1] if "." in nome_arquivo else ""
    if extensao == "xlsx":
        linhas = _ler_xlsx(conteudo)
    elif extensao == "csv":
        linhas = _ler_csv(conteudo)
    else:
        raise ErroDeNegocio("Formato não suportado. Envie um arquivo .xlsx ou .csv.")

    if not linhas:
        raise ErroDeNegocio("A planilha não tem cabeçalho.")
    cabecalho = [_chave_coluna(c) for c in linhas[0]]
    resultado = []
    for numero, valores in enumerate(linhas[1:], start=2):
        celulas = ["" if v is None else str(v).strip() for v in valores]
        if not any(celulas):
            continue  # linha em branco
        resultado.append(LinhaImportada(numero, dict(zip(cabecalho, celulas))))
    if len(resultado) > LINHAS_MAXIMAS:
        raise ErroDeNegocio(f"A planilha tem mais de {LINHAS_MAXIMAS} linhas. Divida em arquivos menores.")
    return resultado


def _ler_xlsx(conteudo: bytes) -> list[list]:
    try:
        planilha = load_workbook(io.BytesIO(conteudo), read_only=True, data_only=True).active
    except Exception as erro:
        raise ErroDeNegocio("Não foi possível abrir o arquivo Excel. Confira se ele não está corrompido.") from erro
    return [list(linha) for linha in planilha.iter_rows(values_only=True)]


def _ler_csv(conteudo: bytes) -> list[list]:
    try:
        texto = conteudo.decode("utf-8-sig")
    except UnicodeDecodeError:
        texto = conteudo.decode("cp1252", errors="replace")  # CSV salvo pelo Excel antigo
    primeira = texto.splitlines()[0] if texto.strip() else ""
    try:
        dialeto = csv.Sniffer().sniff(primeira, delimiters=",;\t")
    except csv.Error:
        dialeto = csv.excel
    return list(csv.reader(io.StringIO(texto), dialeto))


def _campo(dados: dict, *nomes: str) -> str:
    for nome in nomes:
        if dados.get(nome):
            return dados[nome]
    return ""


# --- Questões --------------------------------------------------------------------

COLUNAS_QUESTAO = [
    "enunciado", "alternativa_a", "alternativa_b", "alternativa_c", "alternativa_d", "alternativa_e",
    "gabarito", "disciplina", "categoria", "dificuldade",
]


def validar_questoes(linhas: list[LinhaImportada], enunciados_existentes: list[str]) -> Previa:
    """RF11: campos obrigatórios, alternativas, gabarito, formato e duplicidades."""
    previa = Previa()
    existentes = {normalizar_texto(e) for e in enunciados_existentes}
    vistos: dict[str, int] = {}

    for item in linhas:
        d = item.dados
        erros = []
        enunciado = _campo(d, "enunciado", "pergunta", "questao")
        alternativas = [_campo(d, f"alternativa_{l}", l) for l in "abcde"]
        while alternativas and not alternativas[-1]:
            alternativas.pop()  # E (e além) opcionais
        gabarito = _campo(d, "gabarito", "correta", "resposta").upper()
        dificuldade = normalizar_texto(_campo(d, "dificuldade")) or None

        if not enunciado:
            erros.append("Enunciado não informado.")
        if len(alternativas) < 4 or any(not a for a in alternativas):
            erros.append("Informe as alternativas A, B, C e D (a E é opcional).")
        elif len({normalizar_texto(a) for a in alternativas}) != len(alternativas):
            erros.append("Há alternativas repetidas.")
        letras_validas = "ABCDE"[: max(len(alternativas), 4)]
        if not gabarito:
            erros.append("Gabarito não informado.")
        elif len(gabarito) != 1 or gabarito not in letras_validas:
            erros.append(f"Gabarito '{gabarito}' inválido: use {', '.join(letras_validas)}.")
        if dificuldade and dificuldade not in DIFICULDADES:
            erros.append("Dificuldade inválida: use fácil, média ou difícil.")

        if enunciado:
            chave = normalizar_texto(enunciado)
            if chave in existentes:
                erros.append("Questão já existe no banco de questões.")
            elif chave in vistos:
                erros.append(f"Questão repetida na planilha (igual à linha {vistos[chave]}).")
            else:
                vistos[chave] = item.linha

        if erros:
            previa.erros.append(ErroLinha(item.linha, erros))
        else:
            previa.validas.append(
                LinhaImportada(
                    item.linha,
                    {
                        "enunciado": enunciado,
                        "alternativas": alternativas,
                        "correta": gabarito,
                        "disciplina": _campo(d, "disciplina") or None,
                        "categoria": _campo(d, "categoria") or None,
                        "dificuldade": DIFICULDADES.get(dificuldade) if dificuldade else None,
                    },
                )
            )
    return previa


# --- Alunos ----------------------------------------------------------------------

COLUNAS_ALUNO = ["nome", "matricula", "email"]


def validar_alunos(linhas: list[LinhaImportada], matriculas_existentes: set[str]) -> Previa:
    previa = Previa()
    vistas: dict[str, int] = {}

    for item in linhas:
        d = item.dados
        erros = []
        nome = _campo(d, "nome", "aluno", "nome_completo")
        matricula = _campo(d, "matricula", "ra", "registro")
        email = _campo(d, "email", "e-mail") or None

        if not nome:
            erros.append("Nome não informado.")
        if not matricula:
            erros.append("Matrícula não informada.")
        elif matricula in matriculas_existentes:
            erros.append("Matrícula já cadastrada nesta turma.")
        elif matricula in vistas:
            erros.append(f"Matrícula repetida na planilha (igual à linha {vistas[matricula]}).")
        else:
            vistas[matricula] = item.linha
        if email and not EMAIL.match(email):
            erros.append(f"E-mail '{email}' inválido.")

        if erros:
            previa.erros.append(ErroLinha(item.linha, erros))
        else:
            previa.validas.append(LinhaImportada(item.linha, {"nome": nome, "matricula": matricula, "email": email}))
    return previa


# --- Modelos para download (RF13) ------------------------------------------------

EXEMPLO_QUESTAO = [
    "Qual linguagem define a estrutura de uma página web?", "HTML", "CSS", "JavaScript", "SQL", "",
    "A", "Desenvolvimento Web", "Fundamentos", "facil",
]
EXEMPLO_ALUNO = ["Maria da Silva", "2026001", "maria@exemplo.com"]


def modelo_csv(colunas: list[str], exemplo: list[str]) -> bytes:
    saida = io.StringIO()
    escritor = csv.writer(saida, delimiter=";")  # ";" abre direto no Excel em português
    escritor.writerow(colunas)
    escritor.writerow(exemplo)
    return saida.getvalue().encode("utf-8-sig")


def modelo_xlsx(colunas: list[str], exemplo: list[str]) -> bytes:
    livro = Workbook()
    planilha = livro.active
    planilha.title = "Importação"
    planilha.append(colunas)
    planilha.append(exemplo)
    for celula in planilha[1]:
        celula.font = Font(bold=True, color="FFFFFF")
        celula.fill = PatternFill("solid", fgColor="305496")
    for coluna in planilha.columns:
        planilha.column_dimensions[coluna[0].column_letter].width = max(14, len(str(coluna[1].value or "")) + 2)
    saida = io.BytesIO()
    livro.save(saida)
    return saida.getvalue()
