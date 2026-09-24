"""Relatório da avaliação em Excel (RF47, RF48).

Abas:
- Resultados: aluno, matrícula, turma, avaliação, versão, resposta de cada questão
  (verde = certa, vermelho = errada, cinza = em branco/anulada), gabarito, acertos, erros e nota.
- Questões: acerto e escolhas por alternativa (RF44, RF45).
- Resumo: média, maior e menor nota, distribuição das notas (RF46).
"""

import io

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from app.models.resultado import Resultado
from app.services.estatisticas import Estatisticas
from app.services.nota import ANULADA

AZUL = PatternFill("solid", fgColor="305496")
VERDE = PatternFill("solid", fgColor="C6EFCE")
VERMELHO = PatternFill("solid", fgColor="FFC7CE")
CINZA = PatternFill("solid", fgColor="EDEDED")
BRANCO_NEGRITO = Font(bold=True, color="FFFFFF")
BORDA = Border(*(Side(style="thin", color="BFBFBF"),) * 4)
CENTRO = Alignment(horizontal="center", vertical="center")


def _cabecalho(planilha, colunas: list[str]) -> None:
    planilha.append(colunas)
    for celula in planilha[planilha.max_row]:
        celula.font = BRANCO_NEGRITO
        celula.fill = AZUL
        celula.alignment = CENTRO
        celula.border = BORDA
    planilha.freeze_panes = f"A{planilha.max_row + 1}"


def _larguras(planilha, minimo: int = 6, maximo: int = 45) -> None:
    for coluna in planilha.columns:
        tamanho = max((len(str(c.value)) for c in coluna if c.value is not None), default=0)
        planilha.column_dimensions[get_column_letter(coluna[0].column)].width = max(minimo, min(maximo, tamanho + 2))


def _texto_resposta(resposta: str | None) -> str:
    if resposta is None:
        return "—"
    return "Anulada" if resposta == ANULADA else resposta


def _aba_resultados(planilha, avaliacao: str, resultados: list[Resultado]) -> None:
    planilha.title = "Resultados"
    quantidade = max((len(r.gabarito) for r in resultados), default=0)
    fixas = ["Aluno", "Matrícula", "Turma", "Avaliação", "Versão"]
    finais = ["Gabarito", "Acertos", "Erros", "Em branco", "Anuladas", "Nota", "Nota máxima", "Origem", "Corrigido em"]
    _cabecalho(planilha, fixas + [f"Q{n}" for n in range(1, quantidade + 1)] + finais)

    for r in resultados:
        respostas = [_texto_resposta(r.respostas.get(n)) for n in range(1, quantidade + 1)]
        gabarito = " ".join(f"{n}-{g}" for n, g in sorted(r.gabarito.items()))
        planilha.append(
            [r.aluno_nome or "(sem aluno)", r.matricula, r.turma_nome, avaliacao, r.versao_nome, *respostas, gabarito,
             r.acertos, r.erros, r.em_branco, r.anuladas, r.nota, r.nota_maxima,
             "Leitura automática" if r.origem == "omr" else "Manual", r.corrigido_em.replace(tzinfo=None)]
        )
        linha = planilha.max_row
        for n in range(1, quantidade + 1):
            celula = planilha.cell(row=linha, column=len(fixas) + n)
            celula.alignment = CENTRO
            resposta = r.respostas.get(n)
            if resposta is None or resposta == ANULADA:
                celula.fill = CINZA
            else:
                celula.fill = VERDE if resposta == r.gabarito.get(n) else VERMELHO
        planilha.cell(row=linha, column=len(fixas) + quantidade + 6).number_format = "0.00"
        planilha.cell(row=linha, column=len(fixas) + quantidade + 9).number_format = "dd/mm/yyyy hh:mm"
    _larguras(planilha)


def _aba_questoes(planilha, est: Estatisticas) -> None:
    letras = sorted({l for q in est.por_questao for l in q.escolhas})
    _cabecalho(
        planilha,
        ["Nº", "Enunciado", "Correta", "Respostas", "% de acerto", *[f"{l} (%)" for l in letras],
         "Em branco", "Anuladas", "Mais escolhida"],
    )
    for q in est.por_questao:
        planilha.append(
            [q.ordem, q.enunciado, q.correta, q.respondentes, q.percentual_acerto / 100,
             *[q.percentuais.get(l, 0) / 100 if l in q.escolhas else None for l in letras],
             q.em_branco, q.anuladas, q.mais_escolhida or "—"]
        )
        linha = planilha.max_row
        for coluna in range(5, 6 + len(letras)):
            planilha.cell(row=linha, column=coluna).number_format = "0.0%"
        if q.mais_escolhida and q.mais_escolhida != q.correta:  # distrator forte: destaca
            planilha.cell(row=linha, column=planilha.max_column).fill = VERMELHO
    _larguras(planilha)


def _aba_resumo(planilha, avaliacao: str, est: Estatisticas) -> None:
    linhas = [
        ("Avaliação", avaliacao),
        ("Provas corrigidas", est.provas_corrigidas),
        ("Nota máxima", est.nota_maxima),
        ("Média", est.media),
        ("Mediana", est.mediana),
        ("Maior nota", est.maior_nota),
        ("Menor nota", est.menor_nota),
        ("Desvio padrão", est.desvio_padrao),
        ("Média de acertos", est.media_acertos),
        ("Média de erros", est.media_erros),
        ("% de acerto geral", (est.percentual_acerto or 0) / 100),
    ]
    for rotulo, valor in linhas:
        planilha.append([rotulo, valor])
        planilha.cell(row=planilha.max_row, column=1).font = Font(bold=True)
    planilha.cell(row=planilha.max_row, column=2).number_format = "0.0%"

    planilha.append([])
    planilha.append(["Distribuição das notas", "Provas"])
    for celula in planilha[planilha.max_row]:
        celula.font, celula.fill = BRANCO_NEGRITO, AZUL
    for faixa in est.distribuicao:
        planilha.append([f"{faixa.de:.1f} a {faixa.ate:.1f}", faixa.quantidade])

    planilha.append([])
    planilha.append(["Versão", "Provas", "Média"])
    for celula in planilha[planilha.max_row]:
        celula.font, celula.fill = BRANCO_NEGRITO, AZUL
    for versao in est.por_versao:
        planilha.append([versao.versao, versao.provas, versao.media])
    _larguras(planilha)


def gerar_relatorio(avaliacao: str, resultados: list[Resultado], estatisticas: Estatisticas) -> bytes:
    livro = Workbook()
    _aba_resultados(livro.active, avaliacao, resultados)
    _aba_questoes(livro.create_sheet("Questões"), estatisticas)
    _aba_resumo(livro.create_sheet("Resumo"), avaliacao, estatisticas)
    saida = io.BytesIO()
    livro.save(saida)
    return saida.getvalue()
