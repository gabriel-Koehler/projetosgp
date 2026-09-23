import io

import pytest
from openpyxl import Workbook, load_workbook

from app.services import importacao
from app.services.errors import ErroDeNegocio

CABECALHO_Q = "enunciado;alternativa_a;alternativa_b;alternativa_c;alternativa_d;gabarito;disciplina;dificuldade"


def csv(*linhas: str, codificacao="utf-8") -> bytes:
    return "\n".join(linhas).encode(codificacao)


def xlsx(linhas: list[list]) -> bytes:
    livro = Workbook()
    for linha in linhas:
        livro.active.append(linha)
    saida = io.BytesIO()
    livro.save(saida)
    return saida.getvalue()


def previa_questoes(conteudo: bytes, nome="q.csv", existentes=()):
    return importacao.validar_questoes(importacao.ler_planilha(conteudo, nome), list(existentes))


# --- Leitura dos arquivos -----------------------------------------------------------


def test_le_csv_com_ponto_e_virgula_e_acentos():
    linhas = importacao.ler_planilha(csv(CABECALHO_Q, "Questão ã;a;b;c;d;A;Física;fácil"), "q.csv")
    assert linhas[0].linha == 2
    assert linhas[0].dados["enunciado"] == "Questão ã"
    assert linhas[0].dados["disciplina"] == "Física"


def test_le_csv_com_virgula_e_codificacao_do_excel_antigo():
    conteudo = csv("enunciado,alternativa_a,alternativa_b,alternativa_c,alternativa_d,gabarito", "Ação,a,b,c,d,B",
                   codificacao="cp1252")
    assert importacao.ler_planilha(conteudo, "q.csv")[0].dados["enunciado"] == "Ação"


def test_le_xlsx_e_ignora_linhas_em_branco():
    conteudo = xlsx([["Enunciado", "Alternativa A", "Alternativa B", "Alternativa C", "Alternativa D", "Gabarito"],
                     ["Q1", "a", "b", "c", "d", "C"], [None] * 6, ["Q2", 1, 2, 3, 4, "D"]])
    linhas = importacao.ler_planilha(conteudo, "banco.XLSX")
    assert [l.linha for l in linhas] == [2, 4]
    assert linhas[1].dados["alternativa_a"] == "1"


@pytest.mark.parametrize(
    ("conteudo", "nome", "mensagem"),
    [
        (b"", "q.csv", "vazio"),
        (b"x", "q.pdf", "Formato"),
        (b"lixo", "q.xlsx", "Excel"),
        (b"x" * (importacao.TAMANHO_MAXIMO + 1), "q.csv", "5 MB"),
    ],
    ids=["vazio", "pdf", "xlsx-corrompido", "grande-demais"],
)
def test_arquivos_invalidos(conteudo, nome, mensagem):
    with pytest.raises(ErroDeNegocio, match=mensagem):
        importacao.ler_planilha(conteudo, nome)


# --- Validação das questões (RF11) --------------------------------------------------


def test_questao_valida_vai_para_a_previa():
    previa = previa_questoes(csv(CABECALHO_Q, "O que é HTML?;Marcação;Estilo;Banco;Rede;a;Web;Média"))
    assert previa.erros == []
    assert previa.validas[0].dados == {
        "enunciado": "O que é HTML?", "alternativas": ["Marcação", "Estilo", "Banco", "Rede"], "correta": "A",
        "disciplina": "Web", "categoria": None, "dificuldade": "media",
    }


def test_quinta_alternativa_opcional():
    cabecalho = CABECALHO_Q.replace("alternativa_d;", "alternativa_d;alternativa_e;")
    previa = previa_questoes(csv(cabecalho, "Q;a;b;c;d;e;E;;"))
    assert previa.validas[0].dados["alternativas"] == ["a", "b", "c", "d", "e"]
    assert previa.validas[0].dados["correta"] == "E"


@pytest.mark.parametrize(
    ("linha", "mensagem"),
    [
        (";a;b;c;d;A;;", "Enunciado não informado"),
        ("Q;a;b;;d;A;;", "alternativas A, B, C e D"),
        ("Q;a;b;c;d;;;", "Gabarito não informado"),
        ("Q;a;b;c;d;E;;", "Gabarito 'E' inválido"),
        ("Q;a;b;c;d;AB;;", "inválido"),
        ("Q;a;b;a;d;A;;", "alternativas repetidas"),
        ("Q;a;b;c;d;A;;muito difícil", "Dificuldade inválida"),
    ],
)
def test_erros_de_validacao(linha, mensagem):
    previa = previa_questoes(csv(CABECALHO_Q, linha))
    assert previa.validas == []
    assert any(mensagem in m for m in previa.erros[0].mensagens)
    assert previa.erros[0].linha == 2


def test_duplicidade_na_planilha_e_no_banco():
    previa = previa_questoes(
        csv(CABECALHO_Q, "O que é HTML?;a;b;c;d;A;;", "o que  é  html?;a;b;c;d;A;;", "Já cadastrada;a;b;c;d;A;;"),
        existentes=["JÁ CADASTRADA"],
    )
    assert [v.linha for v in previa.validas] == [2]
    assert previa.erros[0].mensagens == ["Questão repetida na planilha (igual à linha 2)."]
    assert previa.erros[1].mensagens == ["Questão já existe no banco de questões."]


def test_varios_erros_na_mesma_linha():
    previa = previa_questoes(csv(CABECALHO_Q, ";a;b;c;;;;"))
    assert len(previa.erros[0].mensagens) == 3


# --- Validação dos alunos (RF05) ----------------------------------------------------


def test_alunos():
    linhas = importacao.ler_planilha(
        csv("Nome;Matrícula;E-mail", "Ana;1;ana@x.com", "Bia;2;bia", ";3;", "Caio;1;", "Davi;9;"), "a.csv"
    )
    previa = importacao.validar_alunos(linhas, matriculas_existentes={"9"})

    assert [v.dados for v in previa.validas] == [{"nome": "Ana", "matricula": "1", "email": "ana@x.com"}]
    erros = {e.linha: e.mensagens for e in previa.erros}
    assert erros[3] == ["E-mail 'bia' inválido."]
    assert erros[4] == ["Nome não informado."]
    assert erros[5] == ["Matrícula repetida na planilha (igual à linha 2)."]
    assert erros[6] == ["Matrícula já cadastrada nesta turma."]


# --- Modelos (RF13) ------------------------------------------------------------------


def test_modelos_podem_ser_reimportados():
    for gerar, nome in ((importacao.modelo_csv, "m.csv"), (importacao.modelo_xlsx, "m.xlsx")):
        conteudo = gerar(importacao.COLUNAS_QUESTAO, importacao.EXEMPLO_QUESTAO)
        previa = previa_questoes(conteudo, nome)
        assert previa.erros == [] and len(previa.validas) == 1


def test_modelo_xlsx_tem_cabecalho():
    conteudo = importacao.modelo_xlsx(importacao.COLUNAS_ALUNO, importacao.EXEMPLO_ALUNO)
    planilha = load_workbook(io.BytesIO(conteudo)).active
    assert [c.value for c in planilha[1]] == ["nome", "matricula", "email"]
