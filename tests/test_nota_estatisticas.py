"""Nota, estatísticas e relatório Excel (sem banco)."""

import io
from datetime import datetime, timezone

import pytest
from openpyxl import load_workbook

from app.models.resultado import QuestaoDaVersao, Resultado
from app.services import estatisticas
from app.services.nota import ANULADA, calcular_nota, situacao
from app.services.relatorio_excel import gerar_relatorio

# --- Nota (RF37, RF38) ----------------------------------------------------------------


def test_nota_proporcional_aos_acertos():
    gabarito = {1: "A", 2: "B", 3: "C", 4: "D"}
    nota = calcular_nota({1: "A", 2: "B", 3: "A", 4: None}, gabarito, nota_maxima=10)
    assert (nota.acertos, nota.erros, nota.em_branco, nota.anuladas, nota.nota) == (2, 1, 1, 0, 5.0)


def test_anulada_nao_pontua_e_conta_como_erro():
    nota = calcular_nota({1: ANULADA, 2: "B", 3: "C"}, {1: "A", 2: "B", 3: "C"}, nota_maxima=6)
    assert (nota.acertos, nota.erros, nota.anuladas, nota.nota) == (2, 1, 1, 4.0)


def test_nota_arredondada():
    assert calcular_nota({1: "A", 2: "X", 3: "X"}, {1: "A", 2: "B", 3: "C"}, 10).nota == 3.33


def test_respostas_de_outra_versao():
    with pytest.raises(ValueError):
        calcular_nota({1: "A"}, {1: "A", 2: "B"}, 10)


@pytest.mark.parametrize(
    ("resposta", "esperado"), [("A", "correta"), ("B", "errada"), (None, "em_branco"), (ANULADA, "anulada")]
)
def test_situacao(resposta, esperado):
    assert situacao(resposta, "A") == esperado


# --- Estatísticas (RF44 a RF46) -------------------------------------------------------


def resultado(id_, versao_id, versao_nome, respostas, gabarito, nota):
    acertos = sum(1 for n, g in gabarito.items() if respostas[n] == g)
    return Resultado(
        id=id_, avaliacao_id=1, avaliacao_nome="N1", versao_id=versao_id, versao_nome=versao_nome, aluno_id=id_,
        aluno_nome=f"Aluno {id_}", matricula=str(id_), turma_nome="ES-01", semestre_nome="2026/2",
        respostas=respostas, gabarito=gabarito, acertos=acertos, erros=len(gabarito) - acertos, em_branco=0,
        anuladas=0, nota=nota, nota_maxima=10, origem="omr", imagem_path=None,
        corrigido_em=datetime(2026, 9, 23, 10, 0, tzinfo=timezone.utc),
    )


@pytest.fixture
def cenario():
    """Questões 100 e 200. Versão A: ordem original. Versão B: questões trocadas e alternativas invertidas."""
    questoes = [
        QuestaoDaVersao(1, 1, 100, "Q100", ["A", "B", "C", "D"]),
        QuestaoDaVersao(1, 2, 200, "Q200", ["A", "B", "C", "D"]),
        QuestaoDaVersao(2, 1, 200, "Q200", ["D", "C", "B", "A"]),
        QuestaoDaVersao(2, 2, 100, "Q100", ["D", "C", "B", "A"]),
    ]
    selecionadas = [{"questao_id": 100, "ordem": 1, "correta": "A"}, {"questao_id": 200, "ordem": 2, "correta": "B"}]
    gab_a, gab_b = {1: "A", 2: "B"}, {1: "C", 2: "D"}  # mesmas respostas certas, em posições diferentes
    resultados = [
        resultado(1, 1, "A", {1: "A", 2: "B"}, gab_a, 10),  # acertou tudo
        resultado(2, 2, "B", {1: "C", 2: "D"}, gab_b, 10),  # acertou tudo (na versão B)
        resultado(3, 2, "B", {1: "A", 2: None}, gab_b, 0),  # Q200: marcou A da versão B = D original
    ]
    return resultados, questoes, selecionadas


def test_escolhas_convertidas_para_a_letra_original(cenario):
    est = estatisticas.calcular(*cenario, nota_maxima=10)
    q100, q200 = est.por_questao

    assert q100.escolhas == {"A": 2, "B": 0, "C": 0, "D": 0} and q100.em_branco == 1
    assert q100.acertos == 2 and q100.percentual_acerto == pytest.approx(66.7)
    assert q200.escolhas == {"A": 0, "B": 2, "C": 0, "D": 1}
    assert q200.percentuais["D"] == pytest.approx(33.3)
    assert q200.mais_escolhida == "B"


def test_estatisticas_da_turma(cenario):
    est = estatisticas.calcular(*cenario, nota_maxima=10)
    assert est.provas_corrigidas == 3
    assert (est.media, est.mediana, est.maior_nota, est.menor_nota) == (6.67, 10, 10, 0)
    assert est.percentual_acerto == pytest.approx(66.7)
    assert [f.quantidade for f in est.distribuicao] == [1, 0, 0, 0, 0, 0, 0, 0, 0, 2]
    assert [(v.versao, v.provas, v.media) for v in est.por_versao] == [("A", 1, 10), ("B", 2, 5)]


def test_sem_provas_corrigidas(cenario):
    _, questoes, selecionadas = cenario
    est = estatisticas.calcular([], questoes, selecionadas, nota_maxima=10)
    assert est.media is None and est.distribuicao == []
    assert est.por_questao[0].mais_escolhida is None


# --- Relatório Excel (RF47, RF48) -----------------------------------------------------


def test_relatorio_excel(cenario):
    resultados = cenario[0]
    livro = load_workbook(io.BytesIO(gerar_relatorio("N1", resultados, estatisticas.calcular(*cenario, 10))))
    assert livro.sheetnames == ["Resultados", "Questões", "Resumo"]

    aba = livro["Resultados"]
    cabecalho = [c.value for c in aba[1]]
    assert cabecalho[:7] == ["Aluno", "Matrícula", "Turma", "Avaliação", "Versão", "Q1", "Q2"]
    assert {"Gabarito", "Acertos", "Erros", "Nota"} <= set(cabecalho)
    linha = [c.value for c in aba[4]]
    assert linha[:7] == ["Aluno 3", "3", "ES-01", "N1", "B", "A", "—"]
    assert linha[cabecalho.index("Gabarito")] == "1-C 2-D"
    assert linha[cabecalho.index("Nota")] == 0
    assert aba.cell(row=2, column=6).fill.fgColor.rgb.endswith("C6EFCE")  # certa = verde
    assert aba.cell(row=4, column=6).fill.fgColor.rgb.endswith("FFC7CE")  # errada = vermelho

    resumo = {l[0].value: l[1].value for l in livro["Resumo"].iter_rows() if l[0].value}
    assert resumo["Provas corrigidas"] == 3 and resumo["Média"] == 6.67
