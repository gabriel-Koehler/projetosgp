"""Correção completa pela API: folha fotografada -> nota -> resultados -> estatísticas -> Excel."""

import io

import cv2
import numpy as np
import pytest
from omr_utils import fotografar, jpg, pintar
from openpyxl import load_workbook


@pytest.fixture
def prova(api):
    semestre = api.post("/api/semestres", json={"nome": "2026/2"}).json()
    turma = api.post("/api/turmas", json={"semestre_id": semestre["id"], "nome": "ES-01"}).json()
    alunos = [
        api.post(f"/api/turmas/{turma['id']}/alunos", json={"nome": nome, "matricula": str(i)}).json()
        for i, nome in enumerate(["Ana", "Bruno", "Carla"], start=1)
    ]
    ids = [
        api.post(
            "/api/questoes", json={"enunciado": f"Q{i}", "alternativas": ["a", "b", "c", "d"], "correta": "ABCD"[i % 4]}
        ).json()["id"]
        for i in range(1, 6)
    ]
    avaliacao = api.post("/api/avaliacoes", json={
        "nome": "N1 — Engenharia de Software", "turma_id": turma["id"], "questao_ids": ids,
        "configuracao": {"quantidade": 2, "embaralhar_questoes": True, "embaralhar_alternativas": True},
    }).json()
    return {"avaliacao": avaliacao, "alunos": alunos, "turma": turma}


def foto_da_folha(api, avaliacao, versao, marcas, intensidade=None, seed=0) -> bytes:
    resposta = api.get(f"/api/avaliacoes/{avaliacao['id']}/versoes/{versao['codigo']}/folha.png")
    folha = cv2.imdecode(np.frombuffer(resposta.content, np.uint8), cv2.IMREAD_GRAYSCALE)
    alternativas = [len(q["alternativas"]) for q in versao["questoes"]]
    return jpg(fotografar(pintar(folha, alternativas, marcas, intensidade), seed=seed))


def corrigir(api, foto: bytes, aluno_id=None):
    dados = {"aluno_id": str(aluno_id)} if aluno_id else {}
    return api.post("/api/correcoes", files={"imagem": ("folha.jpg", foto, "image/jpeg")}, data=dados)


def gabarito(versao) -> dict[int, str]:
    return {int(n): l for n, l in versao["gabarito"].items()}


def test_fluxo_completo_de_correcao(api, prova):
    avaliacao, (ana, bruno, carla) = prova["avaliacao"], prova["alunos"]
    va, vb = avaliacao["versoes"]
    gab_a, gab_b = gabarito(va), gabarito(vb)

    # Ana (versão A) acerta tudo; Bruno (versão B) erra a 1 e deixa a 2 em branco.
    errada = next(l for l in "ABCD" if l != gab_b[1])
    marcas_bruno = {**gab_b, 1: errada}
    del marcas_bruno[2]

    r_ana = corrigir(api, foto_da_folha(api, avaliacao, va, gab_a), ana["id"])
    assert r_ana.status_code == 201, r_ana.text
    r_bruno = corrigir(api, foto_da_folha(api, avaliacao, vb, marcas_bruno, seed=3), bruno["id"]).json()

    resultado_ana = r_ana.json()["resultado"]
    assert (resultado_ana["acertos"], resultado_ana["nota"], resultado_ana["versao_nome"]) == (5, 10.0, "A")
    assert resultado_ana["aluno_nome"] == "Ana" and resultado_ana["origem"] == "omr"
    assert r_bruno["resultado"]["nota"] == 6.0
    assert r_bruno["resultado"]["em_branco"] == 1 and r_bruno["resultado"]["erros"] == 1
    q1 = r_bruno["resultado"]["questoes"][0]
    assert q1 == {"numero": 1, "marcada": errada, "correta": gab_b[1], "situacao": "errada"}

    # Carla: o professor registra manualmente (ex.: folha rasgada).
    manual = api.post("/api/correcoes/manual", json={
        "codigo": va["codigo"], "aluno_id": carla["id"],
        "respostas": {str(n): ("*" if n == 1 else None) for n in gab_a},
    })
    assert manual.status_code == 201, manual.text
    assert manual.json()["anuladas"] == 1 and manual.json()["nota"] == 0

    # RF41: consulta dos resultados.
    lista = api.get(f"/api/avaliacoes/{avaliacao['id']}/resultados").json()
    assert [r["aluno_nome"] for r in lista] == ["Ana", "Bruno", "Carla"]
    assert api.get(f"/api/resultados/{resultado_ana['id']}").json() == lista[0]

    # RF44 a RF46: estatísticas convertidas para as questões originais.
    est = api.get(f"/api/avaliacoes/{avaliacao['id']}/estatisticas").json()
    assert est["provas_corrigidas"] == 3
    assert est["media"] == pytest.approx((10 + 6 + 0) / 3, abs=0.01)
    assert len(est["por_questao"]) == 5
    assert all(q["respondentes"] == 3 for q in est["por_questao"])
    assert sum(q["acertos"] for q in est["por_questao"]) == 5 + 3

    # RF48: exportação Excel.
    excel = api.get(f"/api/avaliacoes/{avaliacao['id']}/resultados.xlsx")
    assert excel.status_code == 200
    assert "resultados_N1_Engenharia_de_Software.xlsx" in excel.headers["content-disposition"]
    aba = load_workbook(io.BytesIO(excel.content))["Resultados"]
    assert [aba.cell(row=i, column=1).value for i in (2, 3, 4)] == ["Ana", "Bruno", "Carla"]

    # Depois de corrigir, a avaliação não pode mais ser excluída.
    assert api.delete(f"/api/avaliacoes/{avaliacao['id']}").status_code == 409


def test_leitura_duvidosa_nao_grava(api, prova):
    avaliacao, aluno = prova["avaliacao"], prova["alunos"][0]
    versao = avaliacao["versoes"][0]
    foto = foto_da_folha(api, avaliacao, versao, gabarito(versao), intensidade={3: 200})
    resposta = corrigir(api, foto, aluno["id"])

    assert resposta.status_code == 422
    corpo = resposta.json()
    assert corpo["codigo"] == "leitura_duvidosa"
    assert corpo["ilegiveis"] == [3]
    assert corpo["codigo_versao"] == versao["codigo"]
    assert "Nada foi registrado" in corpo["detail"]
    assert api.get(f"/api/avaliacoes/{avaliacao['id']}/resultados").json() == []


def test_corrigir_de_novo_substitui_o_resultado_do_aluno(api, prova):
    avaliacao, aluno = prova["avaliacao"], prova["alunos"][0]
    versao = avaliacao["versoes"][0]
    base = {"codigo": versao["codigo"], "aluno_id": aluno["id"]}
    api.post("/api/correcoes/manual", json={**base, "respostas": {n: None for n in versao["gabarito"]}})
    api.post("/api/correcoes/manual", json={**base, "respostas": versao["gabarito"]})

    lista = api.get(f"/api/avaliacoes/{avaliacao['id']}/resultados").json()
    assert len(lista) == 1 and lista[0]["nota"] == 10


def test_folha_sem_aluno(api, prova):
    avaliacao = prova["avaliacao"]
    versao = avaliacao["versoes"][1]
    resposta = corrigir(api, foto_da_folha(api, avaliacao, versao, gabarito(versao)))
    assert resposta.status_code == 201
    assert resposta.json()["resultado"]["aluno_id"] is None


def test_aluno_de_outra_turma(api, prova):
    semestre = api.post("/api/semestres", json={"nome": "2027/1"}).json()
    turma = api.post("/api/turmas", json={"semestre_id": semestre["id"], "nome": "Outra"}).json()
    intruso = api.post(f"/api/turmas/{turma['id']}/alunos", json={"nome": "X", "matricula": "99"}).json()
    versao = prova["avaliacao"]["versoes"][0]
    resposta = api.post("/api/correcoes/manual", json={
        "codigo": versao["codigo"], "aluno_id": intruso["id"], "respostas": versao["gabarito"],
    })
    assert resposta.status_code == 422
    assert "não é da turma" in resposta.json()["detail"]


@pytest.mark.parametrize(
    ("respostas", "trecho"),
    [({"1": "A"}, "Informe a resposta das 5"), ({"1": "Z", "2": None, "3": None, "4": None, "5": None}, "inválida")],
)
def test_correcao_manual_invalida(api, prova, respostas, trecho):
    versao = prova["avaliacao"]["versoes"][0]
    resposta = api.post("/api/correcoes/manual", json={"codigo": versao["codigo"], "respostas": respostas})
    assert resposta.status_code == 422
    assert trecho in resposta.json()["detail"]


def test_excluir_resultado(api, prova):
    versao = prova["avaliacao"]["versoes"][0]
    resultado = api.post("/api/correcoes/manual", json={"codigo": versao["codigo"], "respostas": versao["gabarito"]})
    assert api.delete(f"/api/resultados/{resultado.json()['id']}").status_code == 204
    assert api.get(f"/api/resultados/{resultado.json()['id']}").status_code == 404
